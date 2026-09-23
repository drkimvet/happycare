#!/usr/bin/env python3
"""Build DVM 612 Week 7 lecture: suture patterns and surgical closure.

Lewyt College of Veterinary Medicine (LIU)
Aligned to the public DVM 612 outline (Week 7) and course LOs 1, 3, 9.
Required texts: Fossum 2018; Hendrickson & Baird 2013 (named only).
Week 8 owns materials, needles, and knot tables. Week 13 owns field LA.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x2C, 0x4A)
GOLD = RGBColor(0xC5, 0xA3, 0x5A)
GOLD_LT = RGBColor(0xF4, 0xEB, 0xD3)
RED = RGBColor(0x8B, 0x2E, 0x2E)
RED_LT = RGBColor(0xF8, 0xE8, 0xE8)
GREEN = RGBColor(0x2E, 0x6B, 0x4F)
GREEN_LT = RGBColor(0xE6, 0xF3, 0xEC)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
OFFWHITE = RGBColor(0xF6, 0xF7, 0xF9)
INK = RGBColor(0x1C, 0x1C, 0x1C)
SLATE = RGBColor(0x3D, 0x4A, 0x57)

W, H = Inches(13.333), Inches(7.5)
FOOTER = "DVM 612  |  Dr. Yujin Kim, D.V.M., Ph.D., FFCP  |  Lewyt CVM"
EXPECTED_SLIDES = 40

PT_FOOTER = 12
PT_TITLE = 30
PT_BODY = 18
PT_CARD = 16
PT_CARD_TITLE = 18
PT_SMALL = 14
PT_REF = 13
MIN_PT = 12

OUT_DIR = Path(__file__).resolve().parent
PPTX = OUT_DIR / "DVM-612_Week7_Suture_Patterns_Closures.pptx"

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]


def _solid(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_rect(slide, l, t, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    _solid(sh, color)
    return sh


def add_round(slide, l, t, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    _solid(sh, color)
    try:
        sh.adjustments[0] = 0.08
    except Exception:
        pass
    return sh


def add_line(slide, l, t, w, h, color=NAVY, weight=1.75):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    _solid(sh, color)
    return sh


def set_tf(tf, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set(
            "anchor",
            {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}.get(anchor, "t"),
        )
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    return p


def add_text(slide, l, t, w, h, text, size=18, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Calibri", anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(l, t, w, h)
    set_tf(box.text_frame, text, size=size, bold=bold, color=color, align=align, font=font, anchor=anchor)
    return box


def add_bullets(slide, l, t, w, h, items, size=18, spacing=8):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(spacing)
        run = p.add_run()
        run.text = "▸  " + item
        run.font.size = Pt(size)
        run.font.color.rgb = INK
        run.font.name = "Calibri"
    return box


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def footer_bar(slide, num, total):
    add_rect(slide, 0, Inches(7.18), W, Inches(0.32), NAVY)
    add_text(
        slide,
        Inches(0.35),
        Inches(7.18),
        Inches(10.8),
        Inches(0.32),
        FOOTER,
        size=PT_FOOTER,
        color=GOLD_LT,
        anchor=MSO_ANCHOR.MIDDLE,
    )
    add_text(
        slide,
        Inches(11.2),
        Inches(7.18),
        Inches(1.8),
        Inches(0.32),
        f"{num}  /  {total}",
        size=PT_FOOTER,
        color=WHITE,
        align=PP_ALIGN.RIGHT,
        anchor=MSO_ANCHOR.MIDDLE,
    )


def header_bar(slide):
    add_rect(slide, 0, 0, W, Inches(0.12), GOLD)
    add_rect(slide, 0, Inches(0.12), W, Inches(0.08), NAVY)


def card(slide, l, t, w, h, title, body, fill=WHITE, title_color=NAVY, accent=GOLD):
    add_round(slide, l, t, w, h, fill)
    add_rect(slide, l, t, Inches(0.10), h, accent)
    add_text(slide, l + Inches(0.28), t + Inches(0.12), w - Inches(0.4), Inches(0.36), title, size=PT_CARD_TITLE, bold=True, color=title_color)
    add_text(slide, l + Inches(0.28), t + Inches(0.50), w - Inches(0.4), h - Inches(0.62), body, size=PT_CARD, color=SLATE)


def pill(slide, l, t, w, h, text, fill=GOLD, text_color=NAVY):
    add_round(slide, l, t, w, h, fill)
    add_text(slide, l, t, w, h, text, size=PT_SMALL, bold=True, color=text_color, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def new_content(title):
    s = prs.slides.add_slide(BLANK)
    add_rect(s, 0, 0, W, H, OFFWHITE)
    header_bar(s)
    add_rect(s, 0, Inches(0.12), Inches(0.12), Inches(7.0), GOLD)
    add_text(s, Inches(0.5), Inches(0.50), Inches(12.3), Inches(0.58), title, size=PT_TITLE, bold=True, color=NAVY)
    return s


def new_section(part, title, subtitle, minutes):
    s = prs.slides.add_slide(BLANK)
    add_rect(s, 0, 0, W, H, NAVY)
    add_rect(s, 0, 0, W, Inches(0.14), GOLD)
    add_rect(s, 0, Inches(7.36), W, Inches(0.14), GOLD)
    add_text(s, Inches(0.7), Inches(2.15), Inches(12), Inches(0.4), part.upper(), size=PT_SMALL, bold=True, color=GOLD)
    add_text(s, Inches(0.7), Inches(2.55), Inches(12), Inches(1.1), title, size=32, bold=True, color=WHITE)
    add_text(s, Inches(0.7), Inches(3.75), Inches(12), Inches(0.7), subtitle, size=PT_BODY, color=GOLD_LT)
    pill(s, Inches(0.7), Inches(4.7), Inches(2.4), Inches(0.42), minutes, fill=GOLD, text_color=NAVY)
    return s


def stamp_footers():
    total = len(prs.slides)
    for i, slide in enumerate(prs.slides, 1):
        dark = False
        try:
            rgb = slide.shapes[0].fill.fore_color.rgb
            dark = rgb == NAVY
        except Exception:
            dark = False
        if dark:
            add_text(
                slide,
                Inches(11.2),
                Inches(7.18),
                Inches(1.7),
                Inches(0.32),
                f"{i}  /  {total}",
                size=PT_FOOTER,
                color=GOLD,
                align=PP_ALIGN.RIGHT,
                anchor=MSO_ANCHOR.MIDDLE,
            )
        else:
            footer_bar(slide, i, total)


def assert_min_font(presentation, min_pt=MIN_PT):
    bad = []
    for i, slide in enumerate(presentation.slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if not r.text.strip() or r.font.size is None:
                        continue
                    pt = r.font.size.pt
                    if pt + 0.05 < min_pt:
                        bad.append(f"slide {i}: {pt:.0f} pt  {r.text[:50]!r}")
    if bad:
        raise SystemExit("Type below hall floor (" + str(min_pt) + " pt):\n" + "\n".join(bad))


def assert_title_hygiene(presentation):
    bad = []
    for i, slide in enumerate(presentation.slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.size and r.font.size.pt >= PT_TITLE - 0.5 and "\u2014" in r.text:
                        bad.append(f"slide {i} title has em dash: {r.text!r}")
    if bad:
        raise SystemExit("Em dash in a title:\n" + "\n".join(bad))


def draw_incision(slide, x, y, length=Inches(3.4)):
    add_line(slide, x, y, length, Inches(0.04), NAVY)


def draw_x(slide, cx, cy, span=Inches(0.22), color=GOLD):
    add_line(slide, cx - span / 2, cy, span, Inches(0.035), color)
    add_line(slide, cx, cy - span / 2, Inches(0.035), span, color)


def pattern_panel(slide, l, t, w, h, title):
    add_round(slide, l, t, w, h, WHITE)
    add_rect(slide, l, t, w, Inches(0.42), NAVY)
    add_text(slide, l, t, w, Inches(0.42), title, size=16, bold=True, color=GOLD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    return l + Inches(0.25), t + Inches(0.70)


# =============================================================================
# SLIDES
# =============================================================================

s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, W, H, NAVY)
add_rect(s, 0, 0, W, Inches(0.16), GOLD)
add_rect(s, Inches(0), Inches(0), Inches(0.22), H, GOLD)
add_text(s, Inches(0.75), Inches(0.95), Inches(12), Inches(0.40), "LONG ISLAND UNIVERSITY  ·  LEWYT COLLEGE OF VETERINARY MEDICINE", size=16, bold=True, color=GOLD)
add_text(s, Inches(0.75), Inches(1.45), Inches(12), Inches(0.40), "DVM 612  ·  PRINCIPLES OF SURGERY  ·  WEEK 7", size=20, bold=True, color=GOLD_LT)
add_text(s, Inches(0.75), Inches(2.05), Inches(12), Inches(1.6), "Suture Patterns\nand Surgical Closure", size=36, bold=True, color=WHITE)
add_text(s, Inches(0.75), Inches(3.85), Inches(12), Inches(0.50), "Dr. Yujin Kim, D.V.M., Ph.D., FFCP", size=24, bold=True, color=GOLD)
add_text(s, Inches(0.75), Inches(4.40), Inches(12), Inches(0.40), "Lecture  |  40 slides  |  50 minutes", size=18, color=GOLD_LT)
add_text(s, Inches(0.75), Inches(5.00), Inches(12), Inches(0.90), "Pick the pattern after you name the tissue, the tension, and the blood supply. Then close in layers.", size=20, color=WHITE)
add_text(s, Inches(0.75), Inches(6.05), Inches(12), Inches(0.70), "Required: Fossum 2018; Hendrickson & Baird 2013.\nWeek 8 owns materials, needles, and knot tables. Week 13 owns large-animal field patterns.", size=16, color=GOLD)
notes(s, "Welcome. This is Week 7 of the public DVM 612 outline: suture patterns and surgical closure. Fifty minutes. Forty slides. Do not dump a Fossum size table. Do not teach milligrams. Week 8 is materials and knots. If time is short, protect interrupted versus continuous, appose versus invert, simple interrupted, vertical mattress far-far near-near, Cushing versus Connell, and layered closure.")

s = new_content("Learning objectives")
add_bullets(s, Inches(0.55), Inches(1.20), Inches(12.2), Inches(4.2), [
    "Name a closure as interrupted or continuous, and as appositional, inverting, or everting.",
    "Choose a pattern after you name the tissue, the tension, and the blood supply (course LO 1, 3).",
    "Close skin, subcutis, and body wall as separate layers. Do not skip a layer because the skin looks closed.",
    "Spot the common mistakes: overtightening, one-knot continuous failure, inverting skin, crushing the needle tip (course LO 9).",
    "Leave suture material, needle type, and throw counts for Week 8. Confirm sizes in Fossum or the lab protocol.",
], size=20, spacing=12)
add_round(s, Inches(0.5), Inches(5.55), Inches(12.3), Inches(1.40), GOLD_LT)
add_text(s, Inches(0.75), Inches(5.70), Inches(11.9), Inches(1.10), "Public DVM 612 outline, Week 7. Course texts are named. Bite-width millimeters printed on some teaching pages stay on those pages.", size=16, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Read the five objectives. This hour is judgment plus names. The lab owns the hands. Do not write 3 to 5 millimeters on the board as a law.")

s = new_content("This hour")
rows = [
    ("0–8 min", "Why the pattern matters", "Halsted. Three questions before you sew."),
    ("8–18 min", "How patterns are classified", "Interrupted vs continuous. Appose, invert, evert."),
    ("18–36 min", "The patterns you must name", "Skin, tension, hollow viscus, intradermal."),
    ("36–46 min", "Layered closure", "Skin, subcutis, fascia or linea. Mistakes."),
    ("46–50 min", "Quiz and sources", "Two cases. Then questions."),
]
add_rect(s, Inches(0.45), Inches(1.15), Inches(12.4), Inches(0.48), NAVY)
add_text(s, Inches(0.60), Inches(1.15), Inches(2.4), Inches(0.48), "Clock", size=16, bold=True, color=GOLD, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(3.10), Inches(1.15), Inches(4.4), Inches(0.48), "Block", size=16, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(7.60), Inches(1.15), Inches(5.0), Inches(0.48), "Protect this", size=16, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
for i, (clk, block, protect) in enumerate(rows):
    y = Inches(1.72) + Inches(i * 0.95)
    add_round(s, Inches(0.45), y, Inches(12.4), Inches(0.86), WHITE)
    add_text(s, Inches(0.60), y, Inches(2.4), Inches(0.86), clk, size=18, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(3.10), y, Inches(4.4), Inches(0.86), block, size=18, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(7.60), y, Inches(5.0), Inches(0.86), protect, size=16, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "If discussion runs, cut the pattern-card extras. Keep the clock blocks.")

s = new_content("What this hour is not")
card(s, Inches(0.45), Inches(1.20), Inches(6.15), Inches(2.50), "Week 8", "Suture materials, needle point and curvature, and how many throws make a square knot. Do not invent a size or a throw table tonight.", accent=GOLD)
card(s, Inches(6.80), Inches(1.20), Inches(6.05), Inches(2.50), "Week 9", "Wound healing stages. Today you only need: apposition heals fastest, inversion of skin delays healing.", accent=GOLD)
card(s, Inches(0.45), Inches(3.90), Inches(6.15), Inches(2.90), "Week 13", "Equine and large-animal field patterns. One sentence today: the same three questions still apply in the field.", accent=NAVY)
card(s, Inches(6.80), Inches(3.90), Inches(6.05), Inches(2.90), "The lab", "Hands, models, then tissue. This lecture names the pattern. The lab proves you can place it.", accent=GREEN)
notes(s, "Do not open a materials catalog. If they ask 2-0 nylon, say Week 8 and Fossum or the lab kit.")

s = new_section("Part 1", "Why the pattern matters", "Halsted still owns the closure.", "8 minutes")
notes(s, "Short section card. Then Halsted.")

s = new_content("Halsted principles that close the wound")
principles = [
    ("Gentle tissue handling", "Crushing forceps and a crushed needle tip are still tissue trauma."),
    ("Meticulous hemostasis", "A hematoma is dead space you sewed shut."),
    ("Preserve blood supply", "Overtightening any pattern kills the edge."),
    ("Strict asepsis", "A glove poke contaminates the needle and the suture. Replace both."),
    ("No tension", "Undermine or move tissue. Do not ask the skin suture to do a flap’s job."),
    ("Accurate apposition", "Like tissue to like tissue. Skin to skin. Fascia to fascia."),
    ("Obliterate dead space", "A subcutaneous layer, a walk, or a drain with a pull plan. Not a tighter skin suture."),
]
for i, (t, d) in enumerate(principles):
    col = i % 2
    row = i // 2
    x = Inches(0.5) + Inches(col * 6.4)
    y = Inches(1.18) + Inches(row * 1.42)
    w = Inches(12.3) if i == 6 else Inches(6.2)
    if i == 6:
        x = Inches(0.5)
    add_round(s, x, y, w, Inches(1.32), WHITE)
    add_rect(s, x, y, Inches(0.10), Inches(1.32), GOLD)
    add_text(s, x + Inches(0.28), y + Inches(0.10), w - Inches(0.4), Inches(0.42), f"{i+1}.  {t}", size=18, bold=True, color=NAVY)
    add_text(s, x + Inches(0.28), y + Inches(0.54), w - Inches(0.4), Inches(0.68), d, size=16, color=SLATE)
notes(s, "Two minutes. Every principle has a closure action. Then the three questions.")

s = new_content("Three questions before you sew")
qs = [
    ("1. What tissue is this?", "Skin, subcutis, fascia, hollow viscus, or tendon. The pattern follows the tissue."),
    ("2. What is the tension?", "If the edges gape, fix the tension first. Undermining beats a tighter mattress."),
    ("3. Where is the blood supply?", "Vessels enter perpendicular to the incision. A strangulating pattern delays healing."),
]
for i, (t, d) in enumerate(qs):
    y = Inches(1.20) + Inches(i * 1.75)
    add_round(s, Inches(0.5), y, Inches(12.3), Inches(1.60), WHITE)
    add_rect(s, Inches(0.5), y, Inches(0.12), Inches(1.60), GOLD)
    add_text(s, Inches(0.85), y + Inches(0.18), Inches(11.7), Inches(0.48), t, size=22, bold=True, color=NAVY)
    add_text(s, Inches(0.85), y + Inches(0.72), Inches(11.7), Inches(0.70), d, size=18, color=SLATE)
notes(s, "Board the three questions. They will hear them again on every pattern card.")

s = new_content("Several patterns can be right")
add_bullets(s, Inches(0.55), Inches(1.25), Inches(12.2), Inches(3.6), [
    "More than one pattern can close the same incision (WCVM VSAC Lab 3).",
    "Species, the patient, and the surgeon’s hands all change the choice.",
    "Personal preference is allowed only among appropriate alternatives.",
    "An appropriate alternative still has to answer the three questions.",
], size=20, spacing=14)
add_round(s, Inches(0.5), Inches(5.15), Inches(12.3), Inches(1.75), GOLD_LT)
add_text(s, Inches(0.75), Inches(5.35), Inches(11.9), Inches(1.40), "You do not win by naming the rarest pattern. You win by naming why this one is safe on this tissue tonight.", size=18, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Cite WCVM public lab: several patterns can be appropriate. Do not mock preference. Mock a pattern that inverts skin.")

s = new_section("Part 2", "How patterns are classified", "Interrupted or continuous. Appose, invert, or evert.", "10 minutes")
notes(s, "Classification block.")

s = new_content("Interrupted versus continuous")
card(s, Inches(0.45), Inches(1.20), Inches(6.15), Inches(5.60), "Interrupted", "Each stitch is its own entity.\n\nYou can change tension stitch by stitch.\n\nLoss of one or two stitches usually leaves the line standing.\n\nCosts time, suture, and knots. More foreign material in deep layers.", accent=GOLD)
card(s, Inches(6.80), Inches(1.20), Inches(6.05), Inches(5.60), "Continuous", "Knots only at the start and the end.\n\nFaster. Less suture. Better air and fluid seal.\n\nOne failed knot or one cut strand can open the whole line.\n\nAdjust tension as you go. Do not crush the strand with instruments.", accent=NAVY)
notes(s, "WCVM Lab 3. Continuous fails as a unit. Interrupted fails as a stitch. That is the exam sentence.")

s = new_content("When interrupted is the safer default")
add_bullets(s, Inches(0.55), Inches(1.25), Inches(12.2), Inches(5.4), [
    "Skin on a dog or cat that will swell: you can leave each stitch a little loose.",
    "A contaminated or dirty wound: one stitch can come out without opening the line.",
    "Teaching closures: you can replace one bad stitch.",
    "Uneven tension along the incision: you can tighten only the gaping segment.",
    "If you cannot trust the end knot, do not start a continuous line.",
], size=20, spacing=12)
notes(s, "Default for student skin is interrupted until the lab says otherwise.")

s = new_content("When continuous earns its speed")
add_bullets(s, Inches(0.55), Inches(1.25), Inches(12.2), Inches(5.4), [
    "A long, even, clean layer where you need a seal: subcutis, some hollow organs, some body wall lines.",
    "You must place two secure knots and never nick the standing strand.",
    "Ford interlocking is a continuous compromise: more locks, still one strand (WCVM).",
    "Do not use speed as the reason to close skin you cannot watch for swelling.",
], size=20, spacing=14)
notes(s, "Speed is a benefit, not an indication.")

s = new_content("Appositional, inverting, everting")
cols = [
    (GREEN, "Appositional", "Edges meet anatomically.\n\nFastest healing.\nMost cosmetic.\n\nDefault for skin and for most soft tissue."),
    (GOLD, "Inverting", "Edges turn in, away from you.\n\nUsed when a hollow organ needs a serosal seal.\nNarrows the lumen. Avoid if the lumen is already small.\n\nDo not invert skin."),
    (RED, "Everting", "Edges turn out, toward you.\n\nDelays healing.\nRarely the plan (WCVM).\n\nA vertical mattress can evert. That is a cost, not a goal, on skin."),
]
for i, (c, t, b) in enumerate(cols):
    x = Inches(0.45) + Inches(i * 4.25)
    add_round(s, x, Inches(1.20), Inches(4.05), Inches(5.55), WHITE)
    add_rect(s, x, Inches(1.20), Inches(4.05), Inches(0.62), c)
    add_text(s, x, Inches(1.20), Inches(4.05), Inches(0.62), t, size=20, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, x + Inches(0.22), Inches(2.00), Inches(3.65), Inches(4.50), b, size=18, color=INK)
notes(s, "Board: appose skin. Invert some viscus. Do not evert as a habit.")

s = new_content("Tension is treated, not decorated")
add_bullets(s, Inches(0.55), Inches(1.25), Inches(12.2), Inches(3.8), [
    "If the suture will break or cut through, the tension is too high (WCVM).",
    "Skin tolerates tension poorly. Undermine or move tissue before you pick a ‘tension pattern.’",
    "Tension sutures belong where you cannot recruit more tissue: some fascia, tendon, or ligament closures.",
    "A mattress on ischemic skin is not a reconstructive plan.",
], size=20, spacing=12)
add_round(s, Inches(0.5), Inches(5.25), Inches(12.3), Inches(1.65), RED_LT)
add_text(s, Inches(0.75), Inches(5.45), Inches(11.9), Inches(1.30), "Overtightening any pattern obliterates local blood supply. Dogs and cats swell after surgery. Leave room (WCVM).", size=18, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "This is the blood-supply slide. Say swell out loud.")

s = new_content("Hands before the first bite")
cols = [
    ("Needle holders", "Wide-based tripod grip. Rotate the wrist. Do not push the needle like a thumbtack (WCVM Lab 1 reminder)."),
    ("Tissue forceps", "Pencil grip. Stabilize, then let go. Crushing the dermis makes a bruise that looks like a bad pattern."),
    ("The needle", "Do not hold the tip. Do not hold it with your fingers. A glove poke retires the needle and the suture."),
    ("The line", "Start and finish at or just beyond the corners. Match bite size on both sides. Check that the knot is square."),
]
for i, (t, d) in enumerate(cols):
    col = i % 2
    row = i // 2
    x = Inches(0.45) + Inches(col * 6.45)
    y = Inches(1.20) + Inches(row * 2.80)
    card(s, x, y, Inches(6.25), Inches(2.60), t, d, accent=GOLD)
notes(s, "One minute. Then patterns. No throw table.")

s = new_content("Ligature is not a skin pattern")
add_bullets(s, Inches(0.55), Inches(1.25), Inches(12.2), Inches(5.4), [
    "Course LO 1 names secure ligation with suture patterns. They are different jobs.",
    "A ligature occludes a vessel or a pedicle. A pattern closes a cut.",
    "Circumferential and transfixing ligatures live in Fossum and in the spay/neuter weeks. Do not invent a throw table tonight.",
    "If the pedicle bleeds at 02:00, that is a ligature failure, not a skin-pattern failure.",
    "Week 11 and 12 will use this. Today you only have to stop calling every knot a closure.",
], size=20, spacing=12)
notes(s, "One minute. LO 1. Then pattern cards.")

s = new_section("Part 3", "Patterns you must name", "Skin, tension, hollow viscus, intradermal.", "18 minutes")
notes(s, "Pattern cards. Keep moving.")

s = new_content("Pattern map")
headers = ("Pattern", "Interrupted or continuous", "Edge", "Usual night use")
rows = [
    ("Simple interrupted", "Interrupted", "Appositional", "Skin. Teaching default."),
    ("Simple continuous", "Continuous", "Appositional", "Long even layer. Seal."),
    ("Cruciate", "Interrupted", "Appositional", "Skin. Some tension."),
    ("Horizontal mattress", "Interrupted", "Appose to evert", "Tension. Watch blood supply."),
    ("Vertical mattress", "Interrupted", "Everting", "Tension. Far-far, then near-near."),
    ("Ford interlocking", "Continuous", "Appositional", "Long skin or field line."),
    ("Intradermal", "Continuous", "Appositional", "Dermis. Buried knots."),
    ("Lembert", "Usually continuous", "Inverting", "Hollow viscus. Transverse bites."),
    ("Cushing / Connell", "Continuous", "Inverting", "Hollow viscus. Parallel bites."),
]
add_rect(s, Inches(0.32), Inches(1.10), Inches(12.70), Inches(0.42), NAVY)
add_text(s, Inches(0.40), Inches(1.10), Inches(3.20), Inches(0.42), headers[0], size=14, bold=True, color=GOLD, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(3.65), Inches(1.10), Inches(3.10), Inches(0.42), headers[1], size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(6.80), Inches(1.10), Inches(2.30), Inches(0.42), headers[2], size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
add_text(s, Inches(9.15), Inches(1.10), Inches(3.70), Inches(0.42), headers[3], size=14, bold=True, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)
for i, (a, b, c, d) in enumerate(rows):
    y = Inches(1.56) + Inches(i * 0.58)
    add_round(s, Inches(0.32), y, Inches(12.70), Inches(0.54), WHITE)
    add_text(s, Inches(0.40), y, Inches(3.20), Inches(0.54), a, size=14, bold=True, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(3.65), y, Inches(3.10), Inches(0.54), b, size=14, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(6.80), y, Inches(2.30), Inches(0.54), c, size=14, color=INK, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(9.15), y, Inches(3.70), Inches(0.54), d, size=14, color=SLATE, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "This is the map, not the recipe. Names only. Confirm construction in Fossum and in lab.")

s = new_content("Simple interrupted")
px, py = pattern_panel(s, Inches(0.40), Inches(1.15), Inches(5.70), Inches(5.65), "Schematic  ·  original")
draw_incision(s, px, py + Inches(1.15))
for i in range(4):
    x = px + Inches(0.35) + Inches(i * 0.80)
    draw_x(s, x, py + Inches(1.17))
card(s, Inches(6.30), Inches(1.15), Inches(6.55), Inches(5.65), "What to say", "Interrupted. Appositional. The teaching default for skin.\n\nYou can change tension stitch by stitch. One failed stitch is not a dehiscence.\n\nNot a tension suture. If the edges gape, fix the tension.\n\nLeave room for swelling in dogs and cats.\n\nConfirm bite spacing in Fossum or the lab. Do not invent millimeters.", accent=GOLD)
notes(s, "WCVM: interrupted, appositional, normal tension. Minimal blood-supply hit unless overtightened.")

s = new_content("Simple continuous")
px, py = pattern_panel(s, Inches(0.40), Inches(1.15), Inches(5.70), Inches(5.65), "Schematic  ·  original")
draw_incision(s, px, py + Inches(1.15))
add_line(s, px + Inches(0.20), py + Inches(1.00), Inches(3.10), Inches(0.06), GOLD)
card(s, Inches(6.30), Inches(1.15), Inches(6.55), Inches(5.65), "What to say", "Continuous. Appositional. Faster. Better seal. Less foreign material.\n\nThe whole line fails if a knot or the strand fails.\n\nAdjust after each bite. Do not crush the suture with instruments.\n\nA long clean subcutis or a hollow-organ first layer is a common home. Skin on a swelling patient is not the automatic home.\n\nConfirm construction in Fossum.", accent=NAVY)
notes(s, "Exam trap: one knot failure opens the line.")

s = new_content("Cruciate")
px, py = pattern_panel(s, Inches(0.40), Inches(1.15), Inches(5.70), Inches(5.65), "Schematic  ·  original")
draw_incision(s, px, py + Inches(1.20))
draw_x(s, px + Inches(1.10), py + Inches(1.22), span=Inches(0.55))
draw_x(s, px + Inches(2.20), py + Inches(1.22), span=Inches(0.55))
card(s, Inches(6.30), Inches(1.15), Inches(6.55), Inches(5.65), "What to say", "Interrupted. Appositional. A tension suture (WCVM).\n\nTwo simple interrupted bites share one knot, so the X sits over the incision.\n\nLess blood-supply cost than a horizontal mattress, more than a simple interrupted.\n\nLeave the loop loose enough for swelling.\n\nUseful on some skin closures when you want fewer knots than a row of simples.", accent=GOLD)
notes(s, "Students love this one. Make them leave it loose.")

s = new_content("Horizontal mattress")
card(s, Inches(0.45), Inches(1.20), Inches(6.15), Inches(5.60), "Construction", "Interrupted. Appositional to everting, depending on how hard you pull (WCVM).\n\nThe second bite is parallel to the first. Tension sits lateral to the wound (Merck Professional mattress pages).\n\nMajor blood-supply cost if overtightened.\n\nUseful when you need to spread load. Dangerous on thin or already ischemic skin.", accent=GOLD)
card(s, Inches(6.80), Inches(1.20), Inches(6.05), Inches(5.60), "Do not", "Do not use it to hide a closure that needed undermining.\n\nDo not bury the edge until it blanches.\n\nStents or bolsters are a hospital conversation, not a homemade table.\n\nConfirm bite geometry in Fossum or the lab.", accent=RED)
notes(s, "Merck: mattress transfers tension off the edge. WCVM: major ischemia if tight.")

s = new_content("Vertical mattress")
card(s, Inches(0.45), Inches(1.20), Inches(6.15), Inches(5.60), "Far-far, then near-near", "Interrupted. Everting. A tension suture (WCVM).\n\nMerck Professional: one wide-deep loop, then one narrow-shallow loop. Far-far, then near-near.\n\nAligns deep and superficial tissue in one stitch.\n\nLess blood-supply cost than a horizontal mattress, more than a simple interrupted.", accent=GOLD)
card(s, Inches(6.80), Inches(1.20), Inches(6.05), Inches(5.60), "Do not", "Do not place the near-near first. The shallow loop can cut the skin.\n\nDo not treat eversion as a cosmetic win on dog or cat skin.\n\nCross-hatching and ischemia are the published costs (Merck).\n\nConfirm distances in Fossum or the lab. Do not invent millimeters.", accent=RED)
notes(s, "Say far-far near-near twice. That is the exam phrase.")

s = new_content("Near and far patterns")
add_bullets(s, Inches(0.55), Inches(1.25), Inches(12.2), Inches(5.4), [
    "Interrupted. Appositional. A tension suture (WCVM).",
    "The far bite takes the load. The near bite apposes the edge.",
    "The name is the order of the bites: near-far-far-near, or far-near-near-far.",
    "Less blood-supply cost than a horizontal mattress, more than a simple interrupted.",
    "Use when you need both tension relief and edge-to-edge skin. Still fix the real tension first.",
], size=20, spacing=12)
notes(s, "Do not spend five minutes on the order variants. Name the idea.")

s = new_content("Ford interlocking")
add_bullets(s, Inches(0.55), Inches(1.25), Inches(12.2), Inches(5.4), [
    "Continuous. Appositional. Blood-supply cost similar to simple continuous (WCVM).",
    "Each bite locks. You still have one strand, so one cut can still fail the line.",
    "Adjust tension with each bite. It is hard to fix at the end.",
    "Finish by making a loop and tying, or by a bite through intact skin off the incision (WCVM options A and B).",
    "A compromise between interrupted security and continuous speed. Not a reason to skip undermining.",
], size=20, spacing=12)
notes(s, "Field and large-animal students will see this again in Week 13. Same three questions.")

s = new_content("Intradermal or subcuticular")
card(s, Inches(0.45), Inches(1.20), Inches(6.15), Inches(5.60), "What it is", "Continuous. Appositional. Bites live in the dermis, as close to the surface as you can place them without exiting skin (WCVM).\n\nSuccessive bites backtrack. Buried knots at both ends.\n\nMinimal impact on skin blood supply if you do not strangulate.", accent=GOLD)
card(s, Inches(6.80), Inches(1.20), Inches(6.05), Inches(5.60), "What it is not", "It is not a substitute for a dead-space layer.\n\nIt is not a reason to skip an E-collar conversation.\n\nForceps crushing the dermis leaves a bruise. WCVM: stabilize with thumb and finger, palm the forceps.\n\nBurying the knot is a lab skill (WCVM Part 5). Week 8 owns the knot.", accent=NAVY)
notes(s, "Students will want this on every spay. Make them earn a quiet dermis first.")

s = new_content("Lembert")
add_bullets(s, Inches(0.55), Inches(1.25), Inches(12.2), Inches(5.4), [
    "Usually continuous. Inverting. Bites run perpendicular to the incision: transverse Lembert (WCVM).",
    "On a hollow organ: serosa, muscularis, and submucosa. Not through mucosa.",
    "Farther from the edge means more inversion and a smaller lumen.",
    "More blood-supply cost than a simple continuous.",
    "Confirm the organ and the lumen size before you invert. A small-intestine lumen is not a stomach lumen.",
], size=20, spacing=12)
notes(s, "Do not invent a GI size table. Name layers: no mucosa.")

s = new_content("Cushing versus Connell")
card(s, Inches(0.45), Inches(1.20), Inches(6.15), Inches(5.60), "Cushing  (no L)", "Continuous. Inverting. Bites run parallel to the incision.\n\nDoes not enter the lumen. Serosa, muscularis, submucosa only (WCVM).\n\nMemory: Cushing does not go in.", accent=GOLD)
card(s, Inches(6.80), Inches(1.20), Inches(6.05), Inches(5.60), "Connell  (has L)", "Continuous. Inverting. Bites also run parallel.\n\nDoes enter the lumen (WCVM).\n\nMemory: Connell goes into the lumen. The extra letter is the extra depth.\n\nMore contamination risk. Confirm in Fossum when this is the first layer versus an oversew.", accent=NAVY)
notes(s, "This is the highest-yield viscus pair. Repeat the mnemonic once. Do not invent which organ prefers which.")

s = new_content("Purse-string and other named closures")
add_bullets(s, Inches(0.55), Inches(1.25), Inches(12.2), Inches(5.4), [
    "A purse-string is a circular continuous bite that cinches an opening. Name it. Confirm the organ and the material in Fossum or the lab.",
    "Parker-Kerr and other oversew names live in the textbook. Do not fake a construction on the board.",
    "Staples and glue are hospital tools, not this hour’s pattern list.",
    "If you cannot draw the path of the needle, you do not know the pattern yet.",
], size=20, spacing=12)
notes(s, "Do not perform a purse-string demo with invented bite counts.")

s = new_section("Part 4", "Layered closure", "Like tissue to like tissue. Then check the mistakes.", "10 minutes")
notes(s, "Closure block.")

s = new_content("Close in layers")
layers = [
    ("1. Body wall or linea", "Fascia holds the abdomen. Skin does not. A pretty skin line over a weak linea is a hernia waiting."),
    ("2. Subcutis / dead space", "Bring fat to fat. Obliterate dead space. This is not the skin suture pulled harder."),
    ("3. Skin or intradermal", "Appose. Leave room to swell. If the deeper layers are wrong, this layer cannot save the patient."),
]
for i, (t, d) in enumerate(layers):
    y = Inches(1.20) + Inches(i * 1.80)
    add_round(s, Inches(0.5), y, Inches(12.3), Inches(1.65), WHITE)
    add_rect(s, Inches(0.5), y, Inches(0.14), Inches(1.65), GOLD)
    add_text(s, Inches(0.90), y + Inches(0.18), Inches(11.6), Inches(0.45), t, size=22, bold=True, color=NAVY)
    add_text(s, Inches(0.90), y + Inches(0.70), Inches(11.6), Inches(0.75), d, size=18, color=SLATE)
notes(s, "Linea first in the abdomen. Do not invent bite counts or a 2-0 rule.")

s = new_content("Skin, subcutis, hollow viscus")
card(s, Inches(0.45), Inches(1.20), Inches(4.05), Inches(5.60), "Skin", "Default: appositional interrupted.\n\nCruciate if you want fewer knots.\n\nMattress only after you treat tension.\n\nIntradermal if the dermis is quiet and you can bury the knot.\n\nDo not invert.", accent=GOLD)
card(s, Inches(4.65), Inches(1.20), Inches(4.05), Inches(5.60), "Subcutis", "Simple continuous is common for a long even layer.\n\nInterrupted if the space is irregular or dirty.\n\nThis layer is dead space, not cosmetics.", accent=NAVY)
card(s, Inches(8.85), Inches(1.20), Inches(4.00), Inches(5.60), "Hollow viscus", "Appositional closures are used. Inverting closures are used. The organ and the lumen decide.\n\nCushing does not enter the lumen. Connell does.\n\nConfirm the stack in Fossum. Do not invent a two-layer law.", accent=GREEN)
notes(s, "Three columns. Then mistakes.")

s = new_content("Common closure mistakes")
mistakes = [
    ("Overtightening", "Any pattern. Edge blanches. Dogs and cats swell."),
    ("One bad continuous knot", "The entire line is now the knot."),
    ("Inverting skin", "Delays healing. That pattern belonged on viscus."),
    ("Skipping a layer", "Pretty skin, open dead space or a weak linea."),
    ("Crushing the needle tip", "The next bite tears. Replace the needle."),
    ("Sewing tension", "A mattress is not an advancement flap."),
]
for i, (t, d) in enumerate(mistakes):
    col = i % 2
    row = i // 2
    x = Inches(0.45) + Inches(col * 6.45)
    y = Inches(1.18) + Inches(row * 1.85)
    card(s, x, y, Inches(6.25), Inches(1.70), t, d, accent=RED)
notes(s, "Course LO 9. Read six. Then quiz.")

s = new_content("If the line fails after recovery")
card(s, Inches(0.45), Inches(1.20), Inches(6.15), Inches(5.60), "Skin dehiscence", "The skin stitches are gone or the edges have parted.\n\nAsk whether the subcutis and body wall are still closed.\n\nOne missing interrupted stitch is not the same as a continuous line that unzipped.", accent=GOLD)
card(s, Inches(6.80), Inches(1.20), Inches(6.05), Inches(5.60), "Body-wall failure", "A swelling under intact skin can still be a hernia.\n\nSkin cannot hold the abdomen.\n\nThis is why the linea is its own layer, and why Week 4 already called a weak linea a technical problem.", accent=RED)
notes(s, "Connect to Week 4 language without opening that deck. No patient names.")

s = new_content("Practice on a model first")
add_bullets(s, Inches(0.55), Inches(1.25), Inches(12.2), Inches(5.4), [
    "DVM 612 uses models and tissues before live patients (course description).",
    "A useful pad has more than one layer, so you can place subcutis and skin as separate jobs.",
    "Confirm the pad LIU already stocks. Do not treat a catalog as the syllabus.",
    "Match bite size on both sides. Square the knot. Then name the pattern out loud.",
    "If you cannot draw the needle path, you are not ready for tissue.",
], size=20, spacing=12)
notes(s, "Vendor names stay in this script only: Nautilus Surgical is not a suture pad. Makers Nutrition in Commack is supplements. Vetiqo and SurgiReal exist on the open web. Coordinator decides.")

s = new_content("Species and setting")
card(s, Inches(0.45), Inches(1.20), Inches(6.15), Inches(5.60), "Dog and cat", "They swell. Interrupted skin is the student default until the lab says the dermis is quiet enough for continuous or intradermal.\n\nThis course is models and tissues first (DVM 612 description).\n\nConfirm what pad LIU already stocks. Do not invent a brand as the syllabus.", accent=GOLD)
card(s, Inches(6.80), Inches(1.20), Inches(6.05), Inches(5.60), "Other settings", "Week 13 owns equine and large-animal field patterns. Ford interlocking will come back.\n\nThe three questions do not change in a barn.\n\nHendrickson & Baird is named. This hour does not teach a cow closure table.", accent=NAVY)
notes(s, "No large-animal slide deck. One card. Nautilus/Makers research stays in the instructor script, not on the student slide.")

s = new_content("Quiz. Write the pattern")
card(s, Inches(0.45), Inches(1.20), Inches(6.15), Inches(2.70), "Case A", "Canine ventral midline. Skin edges meet without a gap. The patient will swell. Which class of pattern do you start with on skin, and why?", accent=GOLD)
card(s, Inches(6.80), Inches(1.20), Inches(6.05), Inches(2.70), "Case B", "You need one stitch that takes deep load and then apposes the epidermis. Name the mattress and the bite order.", accent=NAVY)
card(s, Inches(0.45), Inches(4.10), Inches(6.15), Inches(2.70), "Case C", "Parallel inverting bites on intestine. One pattern enters the lumen. Which name has the extra letter?", accent=GREEN)
card(s, Inches(6.80), Inches(4.10), Inches(6.05), Inches(2.70), "Case D", "A continuous subcutis line looks perfect. The end knot slips in recovery. What failed, and why was interrupted safer on a teaching case?", accent=RED)
notes(s, "Do not advance until they write. Keys on the next slide.")

s = new_content("Quiz key")
keys = [
    "A. Interrupted, appositional skin. Swelling. One failed stitch does not open the line.",
    "B. Vertical mattress. Far-far, then near-near (Merck). Everting is the cost.",
    "C. Connell enters the lumen. Cushing does not (WCVM).",
    "D. Continuous fails as a unit. The knot was the closure. Interrupted would have left the other stitches.",
]
for i, t in enumerate(keys):
    y = Inches(1.20) + Inches(i * 1.35)
    add_round(s, Inches(0.5), y, Inches(12.3), Inches(1.20), WHITE)
    add_rect(s, Inches(0.5), y, Inches(0.85), Inches(1.20), GOLD)
    add_text(s, Inches(0.5), y, Inches(0.85), Inches(1.20), str(i + 1), size=22, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(1.55), y, Inches(10.9), Inches(1.20), t, size=18, color=INK, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Read the keys. Then references.")

s = new_content("Key points")
pearls = [
    "Ask tissue, tension, and blood supply before you name a pattern.",
    "Interrupted fails as a stitch. Continuous fails as a line.",
    "Appose skin. Invert some hollow viscus. Do not invert skin.",
    "Far-far, then near-near. Connell enters the lumen. Cushing does not.",
    "Close layers. Skin cannot hold a linea. A mattress cannot replace undermining.",
    "Sizes, needles, and throws are Week 8. Confirm them in Fossum or the lab protocol.",
]
for i, t in enumerate(pearls):
    y = Inches(1.15) + Inches(i * 0.92)
    add_round(s, Inches(0.5), y, Inches(12.3), Inches(0.84), WHITE)
    add_rect(s, Inches(0.5), y, Inches(0.85), Inches(0.84), GOLD)
    add_text(s, Inches(0.5), y, Inches(0.85), Inches(0.84), str(i + 1), size=20, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, Inches(1.55), y, Inches(10.9), Inches(0.84), t, size=18, color=INK, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Six lines. Then questions.")

s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, W, H, NAVY)
add_rect(s, 0, 0, W, Inches(0.16), GOLD)
add_rect(s, 0, 0, Inches(0.22), H, GOLD)
add_text(s, Inches(0.55), Inches(0.26), Inches(12.2), Inches(0.28), "DVM 612  ·  PRINCIPLES OF SURGERY  ·  WEEK 7", size=14, bold=True, color=GOLD)
add_text(s, Inches(0.55), Inches(0.52), Inches(12.2), Inches(0.48), "Questions", size=28, bold=True, color=WHITE)
add_text(s, Inches(0.55), Inches(1.02), Inches(12.2), Inches(0.70), "What three questions do you ask before you sew?\nWhich mattress is far-far, then near-near?", size=16, color=GOLD_LT)
add_text(s, Inches(0.55), Inches(1.78), Inches(12.2), Inches(0.32), "References  (course texts and public teaching pages used in this hour)", size=14, bold=True, color=GOLD)
left_refs = (
    "1. Fossum TW. Small Animal Surgery. 5th ed. Elsevier; 2018. ISBN 978-0-323-44344-9. Named. No page dump. No invented size table.\n\n"
    "2. Hendrickson DA, Baird AN. Turner and McIlwraith’s Techniques in Large Animal Surgery. 4th ed. Wiley-Blackwell; 2013. Named. Week 13 owns field patterns.\n\n"
    "3. Long Island University CVM. DVM 612 Principles of Surgery course outline (public). Week 7: suture patterns and surgical closure. Week 8: materials, needles, knots.\n\n"
    "4. University of Saskatchewan WCVM. VSAC Lab 3. Suture patterns; interrupted; continuous. Public teaching page."
)
right_refs = (
    "5. Merck Manual Professional. How to repair a laceration with vertical mattress sutures. Far-far / near-near; tension and ischemia costs.\n\n"
    "6. Merck Manual Professional. How to repair a laceration with horizontal mattress sutures. Load sharing lateral to the wound.\n\n"
    "7. Halsted principles: standard surgical teaching. Applied here to closure, not as a slogan.\n\n"
    "8. Bite-width millimeters on WCVM pages stay on those pages. Confirm in lab."
)
add_text(s, Inches(0.55), Inches(2.16), Inches(6.05), Inches(4.40), left_refs, size=PT_REF, color=WHITE)
add_text(s, Inches(6.75), Inches(2.16), Inches(6.05), Inches(4.40), right_refs, size=PT_REF, color=WHITE)
add_text(s, Inches(0.55), Inches(6.72), Inches(12.2), Inches(0.32), "Dr. Yujin Kim, D.V.M., Ph.D., FFCP  ·  Lewyt CVM  ·  Long Island University", size=14, color=GOLD)
notes(s, "Take questions. If none: three questions are tissue, tension, blood supply. Vertical mattress is far-far then near-near. Dismiss on time.")

stamp_footers()
assert_min_font(prs)
assert_title_hygiene(prs)
if len(prs.slides) != EXPECTED_SLIDES:
    raise SystemExit(f"Expected {EXPECTED_SLIDES} slides, built {len(prs.slides)}")

OUT_DIR.mkdir(parents=True, exist_ok=True)
prs.save(str(PPTX))

script_lines = [
    "DVM 612 Week 7 instructor script",
    "Instructor: Dr. Yujin Kim, D.V.M., Ph.D., FFCP",
    "",
    "40 slides. 50 minutes. Week 7 of the public DVM 612 outline: suture patterns and surgical closure.",
    "Do not alter the Week 4 37-slide deck. Do not invent Fossum sizes, throw counts, or 2-0 nylon as a law.",
    "Week 8 owns materials, needles, and knots. Week 13 owns large-animal field patterns.",
    "Owner names stay off slides. MoMo and Willie stay off this hour.",
    "",
    "MODELS (zip 11577 homework, not a student slide):",
    "DVM 612 uses models and tissues. Confirm with the course coordinator what LIU already stocks before you buy.",
    "Nautilus Surgical (nautilussurgical.com) is a mid-Atlantic microsurgery distributor (optics, chairs, instruments), not a named LIU suture-pad line.",
    "Makers Nutrition (71 Mall Drive, Commack, NY 11725) is a private-label supplement manufacturer, about 25 miles east of Roslyn Heights 11577. Not a suture trainer.",
    "Named veterinary multilayer pads on the open web include Vetiqo and SurgiReal. Prices stay on those catalogs.",
    "Do not put a vendor quote on the student deck.",
    "",
]
for i, slide in enumerate(prs.slides, 1):
    title = ""
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        t = sh.text_frame.text.strip().split("\n")[0].strip()
        if not t or t.startswith("DVM 612  |"):
            continue
        if "/" in t and t.replace(" ", "").replace("/", "").isdigit():
            continue
        if 4 <= len(t) <= 90:
            title = t
            break
    try:
        speak = slide.notes_slide.notes_text_frame.text.strip()
    except Exception:
        speak = ""
    script_lines.append(f"SLIDE {i}. {title or '(visual)'}")
    if speak:
        script_lines.append(f"    SPEAK: {speak}")
    script_lines.append("")

script_path = OUT_DIR / "DVM-612_Week7_Instructor_Script.txt"
script_path.write_text("\n".join(script_lines), encoding="utf-8")
print(f"Saved {PPTX}")
print(f"Script {script_path}")
print(f"Slides {len(prs.slides)}")
