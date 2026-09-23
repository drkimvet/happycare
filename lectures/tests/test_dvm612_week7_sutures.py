#!/usr/bin/env python3
"""Invariants for the DVM 612 Week 7 suture-pattern lecture."""

from __future__ import annotations

import re
import subprocess
import sys
import unittest
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
BUILDER = ROOT / "build_dvm612_week7_sutures.py"
PPTX = ROOT / "DVM-612_Week7_Suture_Patterns_Closures.pptx"
SCRIPT = ROOT / "DVM-612_Week7_Instructor_Script.txt"
FOOTER = "DVM 612  |  Dr. Yujin Kim, D.V.M., Ph.D., FFCP  |  Lewyt CVM"
EXPECTED = 40
MIN_PT = 12
TITLE_PT = 30


def _runs(prs):
    for i, slide in enumerate(prs.slides, 1):
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.text.strip() and r.font.size is not None:
                        yield i, r


class TestWeek7SutureLecture(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        subprocess.run([sys.executable, str(BUILDER)], check=True, cwd=str(ROOT))
        cls.prs = Presentation(str(PPTX))
        cls.script = SCRIPT.read_text(encoding="utf-8")
        cls.blob = "\n".join(r.text for _, r in _runs(cls.prs))

    def test_slide_count_and_footer(self):
        self.assertEqual(len(self.prs.slides), EXPECTED)
        self.assertEqual(self.prs.slide_width, Inches(13.333))
        self.assertEqual(self.prs.slide_height, Inches(7.5))
        self.assertIn(FOOTER, self.blob)
        self.assertIn("WEEK 7", self.blob)

    def test_hall_type_floor(self):
        small = [(i, r.font.size.pt, r.text[:40]) for i, r in _runs(self.prs) if r.font.size.pt + 0.05 < MIN_PT]
        self.assertEqual(small, [])

    def test_no_em_dash_in_titles(self):
        bad = [
            (i, r.text)
            for i, r in _runs(self.prs)
            if r.font.size.pt >= TITLE_PT - 0.5 and "\u2014" in r.text
        ]
        self.assertEqual(bad, [])

    def test_no_teal_kicker_color_on_header_labels(self):
        # Week 4 used teal 13-pt kickers. This hour must not.
        teal = RGB = (0x1B, 0x6B, 0x7A)
        kickers = []
        for i, r in _runs(self.prs):
            if r.font.size.pt != 13:
                continue
            color = r.font.color.rgb
            if color and (color[0], color[1], color[2]) == teal and r.text.isupper():
                kickers.append((i, r.text))
        self.assertEqual(kickers, [])

    def test_cited_scope_and_no_owner_phi(self):
        low = self.blob.lower()
        self.assertIn("fossum", low)
        self.assertIn("wcvm", low)
        self.assertIn("far-far", low)
        self.assertIn("connell", low)
        self.assertIn("cushing", low)
        self.assertIn("week 8", low)
        self.assertNotRegex(self.blob, r"\bMoMo\b")
        self.assertNotRegex(self.blob, r"\bWillie\b")
        self.assertNotIn("2-0 nylon", low)
        self.assertIsNone(re.search(r"\b3–5\s*mm\b", self.blob))
        self.assertIsNone(re.search(r"\b3-5\s*mm\b", self.blob))

    def test_script_keeps_vendor_homework_off_slides(self):
        self.assertIn("Nautilus Surgical", self.script)
        self.assertIn("Makers Nutrition", self.script)
        self.assertIn("11577", self.script)
        self.assertNotIn("Nautilus Surgical", self.blob)
        self.assertNotIn("Makers Nutrition", self.blob)


if __name__ == "__main__":
    unittest.main()
